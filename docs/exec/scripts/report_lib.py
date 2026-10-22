from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = '맑은 고딕'
INK = (31, 41, 55)          # body text
MUTED = (107, 114, 128)     # captions, header/footer
NAVY = (31, 56, 100)        # H1
ACCENT = (46, 84, 150)      # H2, table label
RULE = 'D5DAE3'             # hairline between table rows
HEAD_FILL = 'EEF1F6'        # table header
SOFT_FILL = 'F5F6F8'        # callout
KEY_FILL = 'EAF0FB'         # key summary callout

def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(a), FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)

def shade(cell, fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcpr.append(shd)

def set_borders(table, outer='A9B1C0', inner=RULE, sides=False):
    """Hairline horizontal rules only: no vertical lines, no box."""
    tblpr = table._element.tblPr
    old = tblpr.find(qn('w:tblBorders'))
    if old is not None: tblpr.remove(old)
    b = OxmlElement('w:tblBorders')
    for edge, color, sz in (('top', outer, 8), ('left', None, 0), ('bottom', outer, 8), ('right', None, 0),
                            ('insideH', inner, 4), ('insideV', None, 0)):
        e = OxmlElement(f'w:{edge}')
        if color is None:
            e.set(qn('w:val'), 'nil')
        else:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), color)
        b.append(e)
    tblpr.append(b)

def fix_widths(table, widths):
    """Fixed layout with explicit grid so Word and Pages both honor column widths."""
    tblpr = table._element.tblPr
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblpr.append(lay)
    tw = tblpr.find(qn('w:tblW'))
    if tw is None:
        tw = OxmlElement('w:tblW'); tblpr.append(tw)
    tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(sum(widths) * 567)))
    for col, w in zip(table.columns, widths):
        col.width = Cm(w)

def cell_margins(table, top=50, bottom=50, left=90, right=90):
    tblpr = table._element.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}'); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
    tblpr.append(m)

def add_runs(p, text, size=10, color=None, base_bold=False):
    """**bold** and [[placeholder]] markup."""
    import re
    for part in re.split(r'(\*\*.+?\*\*|\[\[.+?\]\])', text):
        if not part: continue
        if part.startswith('**'):
            r = p.add_run(part[2:-2]); set_font(r, size, True, color)
        elif part.startswith('[['):
            r = p.add_run('[' + part[2:-2] + ']'); set_font(r, size, True, color)
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            r = p.add_run(part); set_font(r, size, base_bold, color)

class Report:
    def __init__(self, header_text):
        self.d = Document()
        sec = self.d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.0); sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
        st = self.d.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10.5); st.font.color.rgb = RGBColor(*INK)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.35
        for name, size, color, before in (('Heading 1', 17, NAVY, 20), ('Heading 2', 13, ACCENT, 14), ('Heading 3', 11, INK, 10)):
            h = self.d.styles[name]; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
            h.font.color.rgb = RGBColor(*color); h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
            h.element.rPr.rFonts.set(qn('w:ascii'), FONT); h.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_runs(hp, header_text, 8.5, MUTED)
        fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(); set_font(r, 9, color=MUTED)
        for t, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
            if t:
                e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._element.append(e)
            else:
                e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._element.append(e)
        self.fig = 0; self.tab = 0

    def h(self, text, level=1):
        p = self.d.add_heading(level=level)
        r = p.add_run(text); set_font(r, None, True)
        return p

    def p(self, text, size=10.5, align=None, after=None, color=None, bold=False):
        para = self.d.add_paragraph(); add_runs(para, text, size, color, bold)
        if align: para.alignment = align
        if after is not None: para.paragraph_format.space_after = Pt(after)
        return para

    def bullets(self, items, size=10.5):
        for it in items:
            para = self.d.add_paragraph(style='List Bullet'); add_runs(para, it, size)
            para.paragraph_format.space_after = Pt(4)

    def table(self, rows, widths, header=True, size=9, fills=None, caption=None, align_center_from=None):
        size = max(size, 9)
        if caption:
            self.tab += 1
            cp = self.d.add_paragraph(); cp.paragraph_format.keep_with_next = True
            cp.paragraph_format.space_before = Pt(6); cp.paragraph_format.space_after = Pt(3)
            r = cp.add_run(f'표 {self.tab}.  '); set_font(r, 9, True, ACCENT)
            add_runs(cp, caption, 9, MUTED)
        t = self.d.add_table(rows=len(rows), cols=len(rows[0]))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; cell_margins(t, 70, 70, 100, 100); set_borders(t)
        fix_widths(t, widths)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                c = t.cell(i, j); c.width = Cm(widths[j])
                para = c.paragraphs[0]; para.paragraph_format.space_after = Pt(0); para.paragraph_format.line_spacing = 1.2
                is_head = header and i == 0
                add_runs(para, str(val), size, INK, is_head or (j == 0 and len(rows[0]) == 2))
                if align_center_from is not None and j >= align_center_from: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if is_head:
                    shade(c, HEAD_FILL)
                    if align_center_from is not None and j >= align_center_from: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.keep_with_next = True   # header row never sits alone at a page bottom
                elif fills and (i, j) in fills: shade(c, fills[(i, j)])
                c.vertical_alignment = 1  # center
        for ri, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            cs = OxmlElement('w:cantSplit'); trpr.append(cs)
            if header and ri == 0:
                th = OxmlElement('w:tblHeader'); trpr.append(th)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)
        return t

    def callout(self, text, fill='FFF4DE', size=10):
        fill = KEY_FILL if fill in ('EAF1FB', 'EAF0FB') else SOFT_FILL
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_margins(t, 120, 120, 180, 180); set_borders(t, outer=fill, inner=fill); fix_widths(t, [17])
        c = t.cell(0, 0); c.width = Cm(17); shade(c, fill)
        lines = text.split('\n')
        add_runs(c.paragraphs[0], lines[0], size, INK)
        for ln in lines[1:]:
            pp = c.add_paragraph(); add_runs(pp, ln, size, INK); pp.paragraph_format.space_after = Pt(1)
        c.paragraphs[0].paragraph_format.space_after = Pt(3)
        self.d.add_paragraph().paragraph_format.space_after = Pt(4)

    def image(self, path, width_cm, caption):
        para = self.d.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True; para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        self.fig += 1
        cp = self.d.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.space_after = Pt(10)
        r = cp.add_run(f'그림 {self.fig}.  '); set_font(r, 9, True, ACCENT)
        add_runs(cp, caption, 9, MUTED)

    def page_break(self):
        self.d.add_page_break()

    def save(self, path):
        self.d.save(str(path))

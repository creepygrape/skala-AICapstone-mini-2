from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent.parent   # docs/exec
SH, DG, LS = BASE / 'shots', BASE / 'diagrams', BASE / 'langsmith'
OUT = BASE / 'G078_문희주_발표.pptx'

FONT, MONO = 'Malgun Gothic', 'Consolas'
ACCENT, ACCENT_SOFT, ACCENT_MID = '2457D6', 'E8EEFC', '8FA8E8'
NAVY, GRAY, LINE, SOFT = '1B2437', '6B7692', 'DDE2EC', 'F3F5F9'
BLUE = GREEN = AMBER = ACCENT      # legacy names: every highlight uses the single accent
TINT_BLUE = TINT_GREEN = TINT_AMBER = TINT_RED = TINT_GRAY = SOFT
INK, MUTED, WHITE = '1B2437', '5B6577', 'FFFFFF'
CODE_BG, CODE_FG, CODE_HL, CODE_DIM = '1E2433', 'E6E9F2', '9DB7FF', '8FA1C7'


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
NOTES = {}


def set_font(run, size=None, bold=None, color=None, name=FONT, italic=None):
    f = run.font
    f.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set('typeface', name)
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color: f.color.rgb = rgb(color)


def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0.0, wrap=True):
    """paras: list of dicts {runs:[(text,{opts})] | text, size, bold, color, align, after}"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get('align', PP_ALIGN.LEFT)
        if 'after' in p: para.space_after = Pt(p['after'])
        if 'line' in p: para.line_spacing = p['line']
        runs = p.get('runs') or [(p['text'], {})]
        for text, o in runs:
            r = para.add_run()
            r.text = text
            set_font(r, o.get('size', p.get('size', 18)), o.get('bold', p.get('bold', False)),
                     o.get('color', p.get('color', INK)), o.get('name', FONT), o.get('italic'))
    return box


def shape(slide, kind, x, y, w, h, fill, line=None, radius=0.06):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(1.25)
    else: s.line.fill.background()
    s.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    return s


def card(slide, x, y, w, h, fill, line=None, radius=0.06):
    return shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, radius)


def pic(slide, path, x, y, w=None, h=None, frame=True):
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    p = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    if frame:
        p.line.color.rgb = rgb('D5DBEA'); p.line.width = Pt(1)
    return p


def new_slide(title=None, dark=False, num=None, notes=''):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    if title:
        tb(s, 0.6, 0.38, 12.1, 0.95, [{'text': title, 'size': 30, 'bold': True, 'color': WHITE if dark else NAVY, 'line': 1.0}],
           anchor=MSO_ANCHOR.MIDDLE)
    n = len(prs.slides)
    tb(s, 0.6, 7.03, 6, 0.3, [{'text': '광주 3반 G078 문희주  ·  WISH FINDER', 'size': 11, 'color': 'AEB6CC' if dark else '8A93A8'}])
    if num != 'none':
        tb(s, 11.9, 7.03, 0.85, 0.3, [{'text': str(n), 'size': 11, 'color': 'AEB6CC' if dark else '8A93A8', 'align': PP_ALIGN.RIGHT}])
    NOTES[n] = notes
    return s


def code_box(slide, x, y, w, h, lines, hl=(), size=13, caption=None):
    card(slide, x, y, w, h, CODE_BG, radius=0.025)
    paras = []
    for i, ln in enumerate(lines):
        color = CODE_HL if i in hl else (CODE_DIM if ln.strip().startswith('#') else CODE_FG)
        paras.append({'text': ln if ln else ' ', 'size': size, 'color': color, 'after': 0, 'line': 1.08,
                      'runs': [(ln if ln else ' ', {'name': MONO, 'color': color, 'size': size})]})
    tb(slide, x + 0.2, y + 0.18, w - 0.4, h - 0.36, paras)
    if caption:
        tb(slide, x, y + h + 0.06, w, 0.3, [{'text': caption, 'size': 11, 'color': MUTED}])


def _bottom_rule(cell, color=LINE):
    """Hairline under each row. Only lnB is written; PowerPoint is strict about tcPr children, and the fill must follow the borders."""
    tcPr = cell._tc.get_or_add_tcPr()
    ln = tcPr.makeelement(qn('a:lnB'), {'w': '9525'})
    sf = ln.makeelement(qn('a:solidFill'), {}); sf.append(sf.makeelement(qn('a:srgbClr'), {'val': color})); ln.append(sf)
    fill = tcPr.find(qn('a:solidFill'))
    if fill is not None:
        fill.addprevious(ln)       # border before fill (schema order)
    else:
        tcPr.append(ln)


def table(slide, x, y, w, cols, rows, size=16, header_fill=SOFT, row_h=0.5, bold_first=False, hl_rows=(), align='center'):
    gt = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = gt.table
    tblPr = gt._element.graphic.graphicData.tbl.tblPr
    tblPr.set('firstRow', '1'); tblPr.set('bandRow', '0')
    for j, cw in enumerate(cols): t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.1)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(header_fill if i == 0 else (ACCENT_SOFT if i in hl_rows else WHITE))
            _bottom_rule(c)
            tf = c.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if (j == 0 or align == 'left') else PP_ALIGN.CENTER
            for k, part in enumerate(val.split('\n')):
                if k: para.add_line_break()
                r = para.add_run(); r.text = part
                set_font(r, size, True if (i == 0 or (bold_first and j == 0)) else False, INK)
    return gt


def chip(slide, x, y, w, h, text, fill, color=WHITE, size=14, bold=True):
    s = card(slide, x, y, w, h, fill, radius=0.3)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    set_font(r, size, bold, color)
    return s


# =========================================================================== 1. cover
n1 = ('안녕하세요, 광주 3반 G078 문희주입니다. 제 과제는 WISH FINDER입니다. '
      '영상 속 장면이 어렴풋이만 기억날 때, 그 장면을 설명하면 영상의 정확한 시점을 찾아 주는 AI Agent입니다.')
s = new_slide(num='none', notes='[10초] ' + n1)
tb(s, 1.2, 1.9, 10.9, 1.3, [{'text': 'WISH FINDER', 'size': 72, 'bold': True, 'color': INK}], anchor=MSO_ANCHOR.MIDDLE)
tb(s, 1.2, 3.3, 10.9, 1.5, [{'text': '기억 속 장면을 찾아 주는', 'size': 32, 'color': MUTED, 'after': 2},
                            {'text': '근거 기반 AI Agent', 'size': 32, 'bold': True, 'color': ACCENT}])
chip(s, 1.2, 5.0, 4.6, 0.66, '광주 3반  G078  문희주', ACCENT, size=22)
tb(s, 1.2, 5.9, 10.9, 0.9, [{'text': 'AI 서비스 Mini Project  ·  과제 2 (신규 Agent)  ·  첫 적용 대상: NCT WISH 영상 3편', 'size': 18, 'color': MUTED}])

# =========================================================================== 2. problem -> solution + tasks
n2 = ('팬들은 "마니또를 뽑다가 자기 자신을 뽑은 장면"처럼 상황만 기억하고 제목이나 대사는 모릅니다. '
      '그래서 장면을 설명하면 원본 자막으로 검증한 영상 시점을 알려 주는 Agent를 만들었습니다. '
      '세부 task는 파싱·청킹, Hybrid 검색, LangGraph, MCP 서버, 웹 화면, 평가입니다. 자막이 영어 번역본이고 화자 표시가 없어 화자는 단정하지 않습니다.')
s = new_slide('장면을 설명하면, 원본 자막으로 검증한 시점을 찾아 준다', notes='[40초] ' + n2)
table(s, 0.6, 1.55, 6.3, [1.3, 5.0],
      [['구분', '내용'],
       ['문제', '제목도 대사도 모르고 "그 장면"만 기억난다'],
       ['해결', '검색 → 원본 cue 대조 → 의미 판정\n→ 영상 시점과 앞뒤 자막'],
       ['제약', '자막은 영어 번역본이고 화자 표시가 없어,\n화자는 단정하지 않는다']],
      size=15, row_h=0.8, bold_first=True, align='left')
pic(s, SH / '03_result_detail.png', 7.2, 1.55, w=5.6)
tb(s, 7.2, 5.05, 5.6, 0.3, [{'text': '결과 화면: 검증된 장면 카드, 영상, 원본 앞뒤 자막', 'size': 12, 'color': MUTED}])
tb(s, 0.6, 5.0, 6, 0.4, [{'text': '세부 task', 'size': 16, 'bold': True, 'color': INK}])
table(s, 0.6, 5.45, 12.2, [2.0, 2.04, 2.04, 2.04, 2.04, 2.04],
      [['파싱·청킹', 'Hybrid 검색', 'LangGraph', 'MCP 서버', 'API·Vue·SSE', '평가·추적'],
       ['cue 3,750 → 청크 869', 'BM25 + bge-m3, RRF', '6개 노드, 재검색 ≤ 2', '읽기 전용 도구 4개', '진행 단계 실시간', 'A/B/C 20문항']],
      size=12, row_h=0.55)

# =========================================================================== 3. results
n3 = ('결과입니다. 평가는 20문항이고, 장면 질문 18개 중 1위로 정답을 맞힌 수는 키워드 검색 A가 13개, Hybrid 검색 B가 15개, 검증과 재검색을 더한 Agent C가 16개입니다. '
      '차이가 작고 한 번만 실행해서 우위는 단정하지 않습니다. 개선한 v3는 화자 질문 판정 기준, 번역 자막 대응, 중복 재검색 제거, 검색어 재작성을 고쳤습니다. '
      '그 결과 화자 질문 처리가 2개에서 3개로 늘고 LLM 호출이 16% 줄었지만, 입력 토큰은 24% 늘고 틀린 장면을 찾음으로 답한 사례가 1건 생겼습니다. '
      '근거 없는 화자 단정은 0건입니다. 이중 첫 검색은 켜도 결과가 같아 꺼 두었고, 청크는 20초와 10초를 유지했습니다. '
      '한계는 후속 질문과 번역 어휘 차이, Multi-Agent 미구현입니다. 이제 구조와 핵심 코드를 설명드리겠습니다.')
s = new_slide('결과: 16/18을 맞혔고, 근거 없는 화자 단정은 0건', notes='[60초] ' + n3)
# left: A/B/C chart (introduced first in the narration)
tb(s, 0.6, 1.5, 5.6, 0.4, [{'text': 'A / B / C 1위 적중 수 (18문항 중)', 'size': 16, 'bold': True, 'color': MUTED}])
cd = CategoryChartData(); cd.categories = ['A 키워드', 'B Hybrid', 'C Agent']; cd.add_series('1위 적중', (13, 15, 16))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.9), Inches(5.7), Inches(3.0), cd)
ch = gf.chart
ch.has_legend = False; ch.has_title = False
ch.font.size = Pt(14); ch.font.name = FONT
pl = ch.plots[0]; pl.gap_width = 60; pl.has_data_labels = True
dl = pl.data_labels; dl.font.size = Pt(18); dl.font.bold = True; dl.font.name = FONT; dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.number_format = '0'; dl.number_format_is_linked = False
for idx, col in enumerate((GRAY, ACCENT_MID, ACCENT)):
    pt = pl.series[0].points[idx]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(col)
va = ch.value_axis; va.maximum_scale = 18; va.minimum_scale = 0; va.major_unit = 6
va.has_major_gridlines = True; va.major_gridlines.format.line.color.rgb = rgb('E3E7F1')
va.format.line.fill.background(); va.tick_labels.font.size = Pt(13); va.tick_labels.font.color.rgb = rgb(MUTED)
ch.category_axis.tick_labels.font.size = Pt(14); ch.category_axis.format.line.color.rgb = rgb('C9D0E0')
tb(s, 0.6, 4.95, 5.6, 0.7, [{'text': 'A 키워드 · B Hybrid(키워드+의미) · C Agent(Hybrid+검증·재검색)', 'size': 12, 'color': MUTED, 'after': 1},
                           {'text': '같은 영어 검색어 조건 · 1~3문항 차이, 1회 실행이라 우위 단정 안 함', 'size': 12, 'color': MUTED}])
# right: v2 -> v3 table + what was improved
tb(s, 6.6, 1.5, 6.2, 0.4, [{'text': '개선 전/후 (OpenAI, 같은 평가셋)', 'size': 16, 'bold': True, 'color': MUTED}])
table(s, 6.6, 1.95, 6.15, [2.95, 1.0, 1.0, 1.2],
      [['지표', 'v2', 'v3', '변화'],
       ['1위 적중 (18문항)', '16', '16', '같음'],
       ['화자 질문 처리 (4문항)', '2', '3', '+1'],
       ['LLM 호출 (20문항)', '50', '42', '−16%'],
       ['평균 지연', '2.88초', '2.57초', '−11%'],
       ['입력 토큰 (C+C-ko)', '56,941', '70,770', '+24%'],
       ['틀린 장면을 "찾음"', '0건', '1건', '+1']],
      size=14, row_h=0.45, hl_rows=(5, 6))
tb(s, 6.6, 5.12, 6.15, 0.6, [{'text': 'v3 개선: 화자 질문 판정 기준 · 번역 자막 대응 · 중복 재검색 제거 · 검색어 재작성', 'size': 12, 'color': MUTED}])
foot = [('0건', ACCENT, SOFT, ['근거 없이 화자를 단정한 경우']),
        ('16 → 16', ACCENT, SOFT, ['이중 첫 검색을 켜도 결과 같음.', '검색만 80 → 98회', '기본 꺼짐으로 설정']),
        ('20초/10초', ACCENT, SOFT, ['청크 길이 비교 후 유지', '1위 적중: 10/5=11 · 20/10=11', '30/15=9 · 40/20=5']),
        ('한계', INK, SOFT, ['후속 질문·번역 어휘 차이', 'Multi-Agent 미구현'])]
for i, (big, col, tint, small) in enumerate(foot):
    x = 0.6 + i * 3.07
    card(s, x, 5.72, 2.95, 1.25, tint)
    tb(s, x + 0.15, 5.74, 2.65, 0.4, [{'text': big, 'size': 20, 'bold': True, 'color': col}], anchor=MSO_ANCHOR.MIDDLE)
    tb(s, x + 0.15, 6.16, 2.65, 0.78, [{'text': t, 'size': 12, 'color': INK, 'line': 1.0, 'after': 0} for t in small])

# =========================================================================== 4. architecture + tech elements
n4 = ('전체 구조입니다. 화면의 질문은 FastAPI를 거쳐 LangGraph Agent로 가고, Agent는 MCP 서버를 통해서만 자막을 조회합니다. '
      '자막 검색은 키워드와 의미 검색을 합친 Hybrid 방식입니다. 이제 핵심 코드 세 가지, Agent Graph와 MCP, RAG를 설명드리겠습니다.')
s = new_slide('LangGraph · MCP · RAG를 하나의 흐름으로 엮었다', notes='[20초] ' + n4)
pic(s, DG / 'architecture.png', 0.5, 1.6, w=7.4, frame=False)
table(s, 8.1, 1.6, 4.75, [1.55, 3.2],
      [['구성', '내용'],
       ['LangGraph', '분류 · 검색 · 검증 · 재검색 흐름'],
       ['MCP', '공식 SDK 2.x, 로컬 stdio, 읽기 전용 도구 4개'],
       ['RAG', 'BM25 + bge-m3 + RRF, 원본 cue 근거'],
       ['의존성 주입', '생성자 주입, LLM을 환경변수로 교체'],
       ['LangSmith', '노드·LLM 호출·토큰 추적'],
       ['모듈화', 'agent · mcp · retrieval · services 분리'],
       ['테스트', '백엔드 178개 · 프런트 43개']],
      size=13, row_h=0.62, bold_first=True, align='left', hl_rows=(1, 2, 3))

# =========================================================================== 5. agent graph
n5 = ('Agent Graph입니다. 코드를 보면 StateGraph에 분류, 검색, 원본 대조, 판정, 재검색, 정리 여섯 노드를 add_node로 등록합니다. '
      '다음 노드는 add_conditional_edges의 조건 함수가 정합니다. 분류에서 자막으로 알 수 없는 요청이면 곧바로 정리로 가고, '
      '판정에서 관련 후보가 없고 시간이 남아 있을 때만 재검색으로 갑니다. MAX_REPLAN_ROUNDS가 2라서 재검색은 최대 두 번입니다. '
      'LLM이 판단하는 곳은 분류, 판정, 검색어 재작성 세 군데뿐이고, 인용과 시각은 코드가 원본 cue에서 만듭니다. Multi-Agent가 아닌 단일 워크플로입니다.')
s = new_slide('Agent Graph: 경로는 코드가, 판단은 LLM이', notes='[65초] ' + n5)
pic(s, DG / 'langgraph_flow.png', 0.5, 1.55, w=6.1, frame=False)
table(s, 0.6, 4.0, 6.1, [2.6, 3.5],
      [['설계 원칙', '내용'],
       ['다음 노드', '조건 함수가 정한다 (add_conditional_edges)'],
       ['LLM 판단', '분류 · 판정 · 검색어 재작성, 세 곳'],
       ['인용·시각', '코드가 원본 cue에서 생성, LLM은 쓰지 않음'],
       ['구조', '단일 Agent 워크플로 (Multi-Agent 미구현)'],
       ['상한', '재검색 2번 · LLM 20초 · 전체 60초']],
      size=13, row_h=0.48, bold_first=True, align='left')
GRAPH_CODE = [
    "graph = StateGraph(FinderState)",
    "for name in ['router', 'search', 'context',",
    "             'verification', 'replan', 'finalize']:",
    "    graph.add_node(name, self._node(name))",
    "",
    "graph.add_conditional_edges('router',",
    "    lambda s: 'finalize' if s['intent'] in",
    "              {'unsupported', 'chat'} else 'search')",
    "graph.add_conditional_edges('verification',",
    "                            self._after_verification)",
    "graph.add_conditional_edges('replan',",
    "    lambda s: 'finalize' if s.get('stop_replan')",
    "              else 'search')",
    "",
    "def _after_verification(self, state):",
    "    return 'replan' if state.get('can_replan') \\",
    "           else 'finalize'",
    "",
    "MAX_REPLAN_ROUNDS = 2",
]
code_box(s, 6.85, 1.5, 6.0, 5.1, GRAPH_CODE, hl={5, 8, 10, 14, 15, 16, 18}, size=13,
         caption='backend/app/agent/graph.py (발췌, 줄 바꿈만 조정)')

# =========================================================================== 6. MCP
n6 = ('MCP입니다. 공식 SDK 2.x의 MCPServer에서 server.tool 데코레이터로 읽기 전용 도구 네 개를 만들고, 로컬 프로세스로 띄워 stdio로 통신합니다. '
      'Agent는 검색과 원본 조회 두 개를 쓰고, 나머지 두 개는 화면용입니다. 핵심은 검색 결과를 그대로 믿지 않는 것입니다. '
      'search_transcripts가 준 청크는 가공된 텍스트라서, 인용과 시각은 get_scene_context가 돌려주는 원본 cue로 다시 확인합니다. '
      '검색은 계산이 무거워 동시에 한 건만 실행하도록 제한했습니다.')
s = new_slide('MCP: 자막은 도구로만 접근하고, 검증은 원본 cue로 다시 한다', notes='[50초] ' + n6)
table(s, 0.6, 1.6, 6.2, [2.9, 3.3],
      [['도구', '쓰는 곳'],
       ['search_transcripts', 'Agent · 검색 단계'],
       ['get_scene_context', 'Agent · 원본 대조'],
       ['get_video_metadata', 'API 서버 · 시작 시'],
       ['get_cue_neighbors', '화면 · 앞뒤 자막']],
      size=16, row_h=0.56, hl_rows=(1, 2), align='left')
table(s, 0.6, 4.65, 6.2, [2.6, 3.6],
      [['원칙', '내용'],
       ['공식 MCP SDK 2.x', '로컬 프로세스, stdio 통신'],
       ['읽기 전용', '파일 경로를 인자로 받지 않는다'],
       ['검색 결과를 믿지 않는다', '인용·시각은 원본 cue로 다시 확인']],
      size=13, row_h=0.5, bold_first=True, align='left')
MCP_CODE = [
    "@server.tool(structured_output=True)",
    "def search_transcripts(",
    "        query: str, mode='hybrid',",
    "        top_k=3, video_ids=None",
    "        ) -> SearchReply:",
    '    """Search registered transcripts.',
    "    Returns cue-linked candidates,",
    '    not answers."""',
    "    with search_gate.acquire():",
    "        _, retriever = resources()",
    "        return retriever.search(",
    "            query, mode, top_k, video_ids)",
    "",
    "# Agent verifies each candidate by",
    "# re-reading the source cues:",
    "@server.tool(structured_output=True)",
    "def get_scene_context(video_id,",
    "        start_seconds, window_seconds=20)",
    "        -> ContextReply:",
]
code_box(s, 7.05, 1.5, 5.8, 5.1, MCP_CODE, hl={0, 8, 15, 16}, size=13,
         caption='backend/app/mcp/server.py (발췌, 타입 표기 단순화)')

# =========================================================================== 7. RAG
n7 = ('RAG입니다. 자막 3,750줄을 영상별로 20초씩, 10초 겹치게 묶어 청크 869개를 만들었습니다. '
      '키워드 검색 상위 20개와 의미 검색 상위 20개를 RRF로 합치는데, 순위가 r인 후보에 1 나누기 60 더하기 r 점수를 주고 더합니다. '
      '청크는 cue 번호를 가지고 있어 결과가 항상 원본 자막 줄로 돌아갑니다. 청크 길이는 10초 5초, 30초 15초, 40초 20초와도 비교했고 20초 10초를 유지했습니다. '
      '가장 크게 배운 것은 같은 질문도 표현에 따라 결과가 달라져서, 한 번 잘 나온 결과를 믿지 말고 측정으로 확인해야 한다는 점입니다. 감사합니다.')
s = new_slide('RAG: 두 검색을 RRF로 합치고 원본 cue로 되돌린다', notes='[55초] ' + n7)
steps = [('자막 cue 3,750줄', TINT_GRAY, GRAY), ('청크 869개  (20초 단위, 10초 겹침)', TINT_BLUE, BLUE),
         ('BM25 상위 20  +  bge-m3 상위 20', TINT_GREEN, GREEN), ('RRF로 순위 합치기', TINT_AMBER, 'B9770E'),
         ('원본 cue로 되돌리기', TINT_GREEN, GREEN)]
for i, (txt, tint, col) in enumerate(steps):
    y = 1.6 + i * 0.98
    card(s, 0.6, y, 5.9, 0.7, tint, line=col)
    tb(s, 0.85, y, 5.5, 0.7, [{'text': txt, 'size': 17, 'bold': True, 'color': NAVY}], anchor=MSO_ANCHOR.MIDDLE)
    if i < len(steps) - 1:
        tb(s, 3.3, y + 0.68, 0.5, 0.3, [{'text': '▼', 'size': 14, 'color': GRAY, 'align': PP_ALIGN.CENTER}])
RAG_CODE = [
    "fused = {}",
    "for rank in ranks:   # BM25 list, dense list",
    "    for position, i in enumerate(rank, 1):",
    "        fused[i] = fused.get(i, 0) \\",
    "                   + 1 / (60 + position)",
    "",
    "for i in sorted(fused,",
    "                key=lambda i: (-fused[i], i)):",
    "    chunk = self.chunks[i]",
    "    ...  # near-duplicate chunks removed",
    "    candidate = Candidate(chunk=chunk,",
    "                          score=fused[i])",
]
code_box(s, 6.85, 1.5, 6.0, 3.55, RAG_CODE, hl={3, 4}, size=13,
         caption='backend/app/retrieval/hybrid.py (발췌, 줄 바꿈만 조정)')
card(s, 6.85, 5.5, 6.0, 1.15, TINT_AMBER)
tb(s, 7.1, 5.55, 5.5, 1.05, [{'text': '점수 = Σ  1 / (60 + 순위)', 'size': 24, 'bold': True, 'color': NAVY, 'after': 2},
                             {'text': '청크 길이 비교(B, 18문항): 10/5=11 · 20/10=11 · 30/15=9 · 40/20=5 → 20/10 유지. 60은 미비교', 'size': 13, 'color': MUTED}],
   anchor=MSO_ANCHOR.MIDDLE)

# =========================================================================== 8. limits + next
# =========================================================================== backup B1: LangSmith
nb1 = '[백업] 질문 대비: LangSmith 추적. 평가 두 번의 합계가 트레이스 160개, 모델 호출 188회, 토큰 135.5K, 비용 $0.30입니다. 재검색이 있으면 지연·LLM 호출·토큰이 대략 두 배가 됩니다.'
s = new_slide('[백업] LangSmith로 노드별 시간·LLM 호출·토큰을 추적했다', notes=nb1)
pic(s, LS / '01_overview.png', 0.5, 1.5, w=6.3)
tb(s, 0.5, 4.3, 6.3, 0.9, [{'text': '평가 2회 합계: 트레이스 160 · 모델 호출 188 · 토큰 135.5K · 비용 $0.30 (LangSmith 계산값)', 'size': 14, 'color': MUTED}])
pic(s, LS / '02_trace_tree.png', 7.05, 1.5, w=3.0)
pic(s, LS / '03_trace_replan.png', 10.3, 1.5, h=3.9)
tb(s, 7.05, 5.6, 5.8, 1.2, [{'text': '왼쪽: 재검색 없는 건 2.01초 · LLM 2회', 'size': 14, 'color': INK, 'after': 2},
                            {'text': '오른쪽: 재검색 1번 건 4.28초 · LLM 4회', 'size': 14, 'color': INK, 'after': 2},
                            {'text': 'MCP 도구 호출은 LangSmith에 따로 잡히지 않는다 (한계)', 'size': 13, 'color': MUTED}])

# =========================================================================== backup B2: DI + tests
nb2 = '[백업] 질문 대비: 의존성 주입과 테스트. FinderAgent는 모든 의존성을 생성자로 받고 테스트에서는 가짜 모델과 가짜 연결을 넣습니다. 모델 호출을 가짜로 대신해 과금이 없고 결과가 항상 같습니다.'
s = new_slide('[백업] 의존성 주입과 테스트', notes=nb2)
DI_CODE = [
    "class FinderAgent:",
    "    def __init__(self, session, metadata, model, *,",
    "                 mode='hybrid',",
    "                 tool_call=call_tool,",
    "                 timeout=60,",
    "                 telemetry=None,",
    "                 dual_query=False, ...):",
    "        self.session = session   # MCP session",
    "        self.model = model       # Ollama | OpenAI",
    "",
    "# production: real session + build_model()",
    "# tests: FakeSession + ScriptedModel",
]
code_box(s, 0.6, 1.5, 6.3, 3.5, DI_CODE, hl={1, 7, 8}, size=13, caption='backend/app/agent/graph.py (발췌)')
tb(s, 0.6, 5.5, 6.3, 1.3, [{'text': 'LLM 교체: WISH_LLM_BACKEND 환경변수 하나', 'size': 16, 'bold': True, 'color': NAVY, 'after': 3},
                           {'text': 'Ollama(개발) ↔ OpenAI(제출), 나머지 코드는 그대로', 'size': 14, 'color': MUTED}])
table(s, 7.2, 1.55, 5.6, [3.9, 1.7],
      [['백엔드 테스트 영역', '개수'], ['Agent 그래프', '55'], ['LLM 교체', '19'], ['자막·검색', '37'],
       ['설정·API·실시간 전송', '33'], ['평가 도구', '32'], ['실제 MCP 통신', '2']],
      size=15, row_h=0.5)
tb(s, 7.2, 5.2, 5.6, 1.3, [{'text': '백엔드 178개 통과 · 프런트 43개 통과', 'size': 15, 'bold': True, 'color': NAVY, 'after': 3},
                           {'text': '모델은 가짜로 대신해 과금 없음, 실제 stdio MCP 서버 테스트 포함', 'size': 13, 'color': MUTED}])

# =========================================================================== backup B3: dual first query
nb3 = '[백업] 질문 대비: 이중 첫 검색. Router 검색어가 번역 어휘와 달라 못 찾던 질문을, 질문 원문도 함께 검색해 보완한 시도입니다. 구현 중 점수 합산 방식에서 정답이 밀리는 문제를 실제 자막으로 확인해 순위 교차 방식으로 바꿨습니다. 정답 평가셋 20문항으로 끈 실행과 켠 실행을 다시 측정했더니 1위 적중이 둘 다 16개로 같았고, 못 찾은 두 문항도 그대로였습니다. 검색만 80회에서 98회로 늘었고 LLM 호출과 토큰은 같아서 기본값은 꺼 두었습니다. 마니또처럼 못 찾던 질문이 찾아진 것은 평가셋 밖 직접 실행에서만 확인했습니다.'
s = new_slide('[백업] 이중 첫 검색: 번역 어휘 차이를 질문 원문 검색으로 보완', notes=nb3)
cols = [('문제', TINT_RED, ACCENT, ['Router 검색어: "manitto"', '자막: "Secret Santa"', '→ 같은 의미인데 못 찾음']),
        ('방법', TINT_BLUE, BLUE, ['Router 검색어 + 질문 원문', '두 검색을 함께 실행하고', '순위를 번갈아 채워 후보 3개']),
        ('확인', TINT_GREEN, GREEN, ['마니또: 못 찾음 → 찾음', '대가: 질문당 약 1~2초 증가', '평가셋 20문항: 16 → 16, 이득 없음'])]
for i, (head, tint, col, lines) in enumerate(cols):
    x = 0.6 + i * 4.1
    card(s, x, 1.6, 3.9, 3.0, tint)
    chip(s, x + 0.25, 1.8, 1.0, 0.45, head, col, size=16)
    tb(s, x + 0.25, 2.45, 3.45, 2.0, [{'text': t, 'size': 16, 'color': INK, 'after': 6} for t in lines])
card(s, 0.6, 4.9, 12.1, 1.7, TINT_AMBER)
tb(s, 0.9, 5.0, 11.5, 1.5, [{'text': '구현 중에 발견한 성능 저하', 'size': 18, 'bold': True, 'color': NAVY, 'after': 4},
                            {'text': '코드 리뷰 반영 중 점수 합산 방식으로 바꿨더니 단위 테스트는 통과했지만 실제 자막에서 정답이 후보 3개 밖으로 밀렸다 '
                                     '("구운 오징어": 45:53 → 42:41). 원인을 확인하고 순위를 번갈아 채우는 방식으로 고쳐 테스트로 고정했다.',
                             'size': 15, 'color': INK, 'line': 1.1}], anchor=MSO_ANCHOR.MIDDLE)

# =========================================================================== backup B4: tech elements vs evaluation + multi-agent
nb4 = '[백업] 질문 대비: 기술 요소와 Multi-Agent. Multi-Agent는 호출·지연·실패 지점이 늘고 실제 원인이 구조가 아니라 어휘였기 때문에 구현하지 않았습니다. 설계안은 Supervisor, Search Agent, Verifier Agent의 허브-스포크 구조입니다.'
s = new_slide('[백업] 기술 요소 정리와 Multi-Agent를 쓰지 않은 이유', notes=nb4)
table(s, 0.6, 1.55, 7.0, [1.9, 5.1],
      [['기술 요소', '이 프로젝트에서'],
       ['Multi-Agent', '미구현 (단일 Agent 워크플로)'],
       ['MCP', '공식 SDK 2.x, 도구 4개, stdio'],
       ['RAG', 'BM25 + bge-m3 + RRF, 원본 cue'],
       ['의존성 주입', '생성자 주입, LLM 환경변수 교체'],
       ['LangSmith', '노드·LLM 호출·토큰 추적'],
       ['모듈화', 'agent / mcp / retrieval / services'],
       ['단위 테스트', '백엔드 178 · 프런트 43']],
      size=15, row_h=0.56, hl_rows=(1,), align='left')
card(s, 7.9, 1.55, 4.9, 4.5, TINT_BLUE)
tb(s, 8.15, 1.7, 4.4, 4.2, [{'text': 'Multi-Agent를 쓰지 않은 이유', 'size': 17, 'bold': True, 'color': NAVY, 'after': 6},
                            {'text': '일이 순서대로 짧고 병렬로 할 일이 없다', 'size': 14, 'color': INK, 'after': 4},
                            {'text': 'Supervisor는 LLM 호출·지연·실패 지점을 늘린다', 'size': 14, 'color': INK, 'after': 4},
                            {'text': '실제 원인은 구조가 아니라 어휘였다', 'size': 14, 'color': INK, 'after': 10},
                            {'text': '설계안 (구현·측정하지 않음)', 'size': 17, 'bold': True, 'color': NAVY, 'after': 6},
                            {'text': 'Supervisor가 Agent를 고르고, Search Agent가 도구를 스스로 고르며, Verifier Agent는 판정만 한다', 'size': 14, 'color': INK}])

# =========================================================================== notes + save
prs.save(str(OUT))

print('saved', OUT, '| slides', len(prs.slides))
